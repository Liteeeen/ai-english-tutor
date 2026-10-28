# -*- coding: utf-8 -*-
"""把 技术文档.md 转成 Daylight 风格的 DOCX（供后续转 PDF）。

运行：<python> md_to_docx.py [输出路径]
默认输出：docs/技术文档.docx（会覆盖旧版）
"""
import pathlib
import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = pathlib.Path(__file__).parent / "技术文档.md"
OUT = pathlib.Path(__file__).parent / "技术文档.docx"

ACCENT = RGBColor(0x50, 0x67, 0x44)   # 深绿
INK = RGBColor(0x24, 0x28, 0x20)      # 主文字
MUTED = RGBColor(0x6b, 0x72, 0x66)    # 次要文字
CODE_BG = "F2F4EF"


def set_font(run, size=11, bold=False, color=INK, mono=False):
    run.font.name = "Consolas" if mono else "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade(paragraph, hex_color):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def add_runs(paragraph, text, base_size=11, color=INK):
    """解析行内的 **加粗** 与 `行内代码`。"""
    parts = re.split(r"(\*\*.*?\*\*|`[^`]*`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=base_size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_font(run, size=base_size - 1, mono=True, color=ACCENT)
        else:
            run = paragraph.add_run(part)
            set_font(run, size=base_size, color=color)


def main():
    text = SRC.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    for section in doc.sections:
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    i = 0
    total = len(lines)
    in_code = False
    code_buf = []

    while i < total:
        line = lines[i]
        stripped = line.strip()

        # ── 代码块 ──
        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_buf = []
            else:
                in_code = False
                p = doc.add_paragraph()
                shade(p, CODE_BG)
                for j, code_line in enumerate(code_buf):
                    if j > 0:
                        p.add_run().add_break()
                    run = p.add_run(code_line if code_line else " ")
                    set_font(run, size=9.5, mono=True)
                p.paragraph_format.space_after = Pt(8)
                code_buf = []
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # ── 标题 ──
        match = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if match:
            level = len(match.group(1))
            title_text = match.group(2)
            sizes = {1: 20, 2: 16, 3: 13.5, 4: 12, 5: 11.5, 6: 11}
            p = doc.add_paragraph()
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(title_text)
            set_font(run, size=sizes.get(level, 11), bold=True,
                     color=ACCENT if level <= 2 else INK)
            p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        # ── 分隔线 ──
        if re.match(r"^---+$", stripped):
            i += 1
            continue

        # ── 表格 ──
        if stripped.startswith("|"):
            table_lines = []
            while i < total and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for table_line in table_lines:
                cells = [c.strip() for c in table_line.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    continue  # 分隔行
                rows.append(cells)
            if rows:
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=ncols)
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(rows):
                    for ci in range(ncols):
                        cell_text = row[ci] if ci < len(row) else ""
                        cell = table.cell(ri, ci)
                        cell.paragraphs[0].text = ""
                        add_runs(cell.paragraphs[0], cell_text, base_size=10)
                        if ri == 0:
                            for run in cell.paragraphs[0].runs:
                                run.font.bold = True
                doc.add_paragraph()
            continue

        # ── 引用 ──
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            add_runs(p, stripped.lstrip(">").strip(), base_size=10.5, color=MUTED)
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # ── 列表 ──
        match = re.match(r"^[-*]\s+(.*)", stripped)
        if match:
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, match.group(1))
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue
        match = re.match(r"^\d+\.\s+(.*)", stripped)
        if match:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, match.group(1))
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # ── 普通段落 ──
        p = doc.add_paragraph()
        add_runs(p, stripped)
        p.paragraph_format.space_after = Pt(5)
        i += 1

    if len(sys.argv) > 1:
        out_path = pathlib.Path(sys.argv[1])
    else:
        out_path = OUT
    doc.save(out_path)
    print(f"[OK] generated {out_path}")


if __name__ == "__main__":
    main()
